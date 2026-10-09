"""Compare two daily progress snapshots and write the result to an Excel workbook.

Snapshots are the CSVs codeorg_colleges.py saves as output/snapshots/progress_<date>.csv.

Usage:
    python progress_report.py     # rebuild the report from the two newest snapshots
"""

import csv
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
SNAPSHOT_DIR = ROOT / "output" / "snapshots"
REPORT_DIR = ROOT / "output" / "reports"

SNAPSHOT_HEADER = ["College Name", "Section Code", "Display Name", "Done", "Status"]

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
TOTAL_FONT = Font(bold=True)
CHANGE_FILLS = {
    "Newly done": PatternFill("solid", fgColor="C6EFCE"),
    "No longer done": PatternFill("solid", fgColor="FFC7CE"),
    "Status changed": PatternFill("solid", fgColor="FFEB9C"),
    "New student": PatternFill("solid", fgColor="DDEBF7"),
    "Removed student": PatternFill("solid", fgColor="D9D9D9"),
}


def snapshot_path(day):
    return SNAPSHOT_DIR / f"progress_{day.isoformat()}.csv"


def write_snapshot(day, rows):
    SNAPSHOT_DIR.mkdir(parents=True, exist_ok=True)
    path = snapshot_path(day)
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(SNAPSHOT_HEADER)
        writer.writerows(rows)
    return path


def read_snapshot(path):
    with path.open(newline="", encoding="utf-8-sig") as file:
        return [tuple(row[column] for column in SNAPSHOT_HEADER) for row in csv.DictReader(file)]


def snapshot_days():
    """Return the dates that have a snapshot, oldest first."""
    days = []
    for path in SNAPSHOT_DIR.glob("progress_*.csv"):
        try:
            days.append(date.fromisoformat(path.stem.removeprefix("progress_")))
        except ValueError:
            continue
    return sorted(days)


def previous_day(day):
    """Return the newest snapshot date before `day`, or None."""
    earlier = [other for other in snapshot_days() if other < day]
    return earlier[-1] if earlier else None


def index_students(rows):
    """Map section code -> {student key: (done, status)}.

    A few sections contain two students with the same display name, so the key
    includes the occurrence number within the section.
    """
    sections = defaultdict(dict)
    seen = Counter()
    for _, code, name, done, status in rows:
        seen[code, name] += 1
        sections[code][name, seen[code, name]] = (done, status)
    return sections


def describe_change(previous, current):
    if previous is None:
        return "New student"
    if current is None:
        return "Removed student"
    if previous[0] != current[0]:
        return "Newly done" if current[0] == "Yes" else "No longer done"
    if previous[1] != current[1]:
        return "Status changed"
    return ""


def compare(today_rows, previous_rows):
    """Return (summary rows, change rows, all-student rows, missing sections) for the workbook."""
    today = index_students(today_rows)
    previous = index_students(previous_rows) if previous_rows is not None else {}
    college_names = {code: name for name, code, *_ in today_rows}

    summary, changes, students = [], [], []
    for code, college in college_names.items():
        current_students = today[code]
        # A section missing from the previous snapshot has nothing to compare against.
        earlier_students = previous.get(code)
        done = sum(1 for done, _ in current_students.values() if done == "Yes")
        total = len(current_students)

        section_changes = Counter()
        for key, current in current_students.items():
            earlier = earlier_students.get(key) if earlier_students is not None else None
            change = describe_change(earlier, current) if earlier_students is not None else ""
            section_changes[change] += 1
            row = (college, code, key[0], *current, *(earlier or ("", "")), change)
            students.append(row)
            if change:
                changes.append(row)
        for key, earlier in (earlier_students or {}).items():
            if key not in current_students:
                section_changes["Removed student"] += 1
                changes.append((college, code, key[0], "", "", *earlier, "Removed student"))

        if earlier_students is None:
            done_before = newly_done = change_in_done = None
        else:
            done_before = sum(1 for done, _ in earlier_students.values() if done == "Yes")
            newly_done = section_changes["Newly done"]
            change_in_done = done - done_before
        summary.append(
            (
                college,
                code,
                total,
                done,
                total - done,
                done / total if total else 0,
                done_before,
                change_in_done,
                newly_done,
                section_changes["New student"] if earlier_students is not None else None,
                section_changes["Removed student"] if earlier_students is not None else None,
            )
        )
    # Sections captured last time that code.org no longer lists (or that failed to load).
    missing = []
    for code, earlier_students in previous.items():
        if code not in today:
            college = next(name for name, other, *_ in previous_rows if other == code)
            done_before = sum(1 for done, _ in earlier_students.values() if done == "Yes")
            missing.append((college, code, len(earlier_students), done_before))
    return summary, changes, students, missing


def add_table(sheet, header, rows, first_row=1):
    """Write a header and rows, then style the header, filter and column widths."""
    for column, title in enumerate(header, start=1):
        cell = sheet.cell(row=first_row, column=column, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for offset, row in enumerate(rows, start=1):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=first_row + offset, column=column, value=value)

    last_column = get_column_letter(len(header))
    sheet.freeze_panes = sheet.cell(row=first_row + 1, column=1)
    sheet.auto_filter.ref = f"A{first_row}:{last_column}{first_row + len(rows)}"
    for column, title in enumerate(header, start=1):
        longest = max([len(str(title))] + [len(str(row[column - 1] or "")) for row in rows])
        sheet.column_dimensions[get_column_letter(column)].width = min(max(longest + 2, 10), 60)


def colour_changes(sheet, change_column, first_row, row_count):
    for row in range(first_row, first_row + row_count):
        fill = CHANGE_FILLS.get(sheet.cell(row=row, column=change_column).value)
        if fill:
            for column in range(1, change_column + 1):
                sheet.cell(row=row, column=column).fill = fill


def build_report(day, today_rows, compared_day=None, previous_rows=None):
    """Write output/reports/<day>.xlsx and return its path."""
    summary, changes, students, missing = compare(today_rows, previous_rows)
    before = f"Done {compared_day}" if compared_day else "Done previous"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet["A1"] = f"code.org progress on {day}"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = (
        f"Compared with {compared_day}" if compared_day else "No earlier snapshot to compare with"
    )
    header = [
        "College Name",
        "Section Code",
        "Students",
        f"Done {day}",
        "Not Done",
        "% Done",
        before,
        "Change in Done",
        "Newly Done",
        "New Students",
        "Removed Students",
    ]
    add_table(sheet, header, summary, first_row=4)
    total_row = 5 + len(summary)
    sheet.cell(row=total_row, column=1, value="Total")
    for column in (3, 4, 5, 7, 8, 9, 10, 11):
        values = [row[column - 1] for row in summary if row[column - 1] is not None]
        if values or column <= 5:
            sheet.cell(row=total_row, column=column, value=sum(values))
    all_students = sum(row[2] for row in summary)
    sheet.cell(
        row=total_row, column=6, value=sum(row[3] for row in summary) / all_students if all_students else 0
    )
    for column in range(1, len(header) + 1):
        sheet.cell(row=total_row, column=column).font = TOTAL_FONT
    for row in range(5, total_row + 1):
        sheet.cell(row=row, column=6).number_format = "0%"

    if missing:
        note_row = total_row + 2
        sheet.cell(
            row=note_row,
            column=1,
            value=f"In the {compared_day} snapshot but not captured on {day} (not in the totals above)",
        ).font = TOTAL_FONT
        for column, title in enumerate(
            ["College Name", "Section Code", "Students", before], start=1
        ):
            sheet.cell(row=note_row + 1, column=column, value=title).font = TOTAL_FONT
        for offset, row in enumerate(missing, start=2):
            for column, value in enumerate(row, start=1):
                sheet.cell(row=note_row + offset, column=column, value=value)

    student_header = [
        "College Name",
        "Section Code",
        "Display Name",
        "Done",
        "Status",
        "Previous Done",
        "Previous Status",
        "Change",
    ]
    change_sheet = workbook.create_sheet("Changes")
    add_table(change_sheet, student_header, changes)
    colour_changes(change_sheet, 8, 2, len(changes))

    student_sheet = workbook.create_sheet("All Students")
    add_table(student_sheet, student_header, students)
    colour_changes(student_sheet, 8, 2, len(students))

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    path = REPORT_DIR / f"{day.isoformat()}.xlsx"
    try:
        workbook.save(path)
    except PermissionError:
        # The workbook is open in Excel; keep the result under another name.
        path = path.with_stem(path.stem + "_new")
        workbook.save(path)
    return path


def build_report_for(day):
    """Build the report for a day that already has a snapshot."""
    compared_day = previous_day(day)
    previous_rows = read_snapshot(snapshot_path(compared_day)) if compared_day else None
    return build_report(day, read_snapshot(snapshot_path(day)), compared_day, previous_rows)


def main():
    days = snapshot_days()
    if not days:
        print(f"No snapshots in {SNAPSHOT_DIR}. Run codeorg_colleges.py first.")
        return 1
    print(f"Saved {build_report_for(days[-1])}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
